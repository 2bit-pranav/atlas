"""The benchmark: fixed tasks, run end to end against a running Atlas.

    python tests/bench.py                      # every case, thinking off and on
    python tests/bench.py merge f1 --thinking on

Every phase of work is judged against this, so "it feels better" becomes a
number. Needs llama-server and the backend running (`atlas start`). The fixtures
are generated here, the first time, into `.data/bench/fixtures/`: two
lab-report PDFs shaped like real ones (a title page of its own, then content
with diagrams), a third with an apparatus section, and a workbook. Generated,
so the benchmark is the same on every machine and contains nobody's documents.

Checks are only what a script can decide for certain: a file exists, both
sources are in a merged PDF, a column holds people rather than teams, no tool
reported a missing image. Facts (who won what) are recorded for a person to
read, never asserted from memory. Each run writes its full trace -- every
reasoning block, tool call and result -- next to the summary, in
`.data/bench/<timestamp>/`.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

import httpx

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / ".data" / "bench" / "fixtures"
TEAMS = {"mercedes", "ferrari", "mclaren", "red bull", "red bull racing", "williams", "aston martin",
         "alpine", "haas", "sauber", "kick sauber", "racing bulls", "rb", "audi", "cadillac"}


@dataclass
class Case:
    name: str
    prompt: str
    uploads: list[str] = field(default_factory=list)
    checks: Callable[["Run"], dict[str, bool]] | None = None


@dataclass
class Run:
    case: str
    thinking: bool
    chat: str
    turn: dict[str, Any]
    workspace: Path

    @property
    def parts(self) -> list[dict[str, Any]]:
        return self.turn["render"]

    @property
    def answer(self) -> str:
        return "".join(p["text"] for p in self.parts if p["kind"] == "text")

    def tools(self, name: str | None = None) -> list[dict[str, Any]]:
        return [p for p in self.parts if p["kind"] == "tool" and (name is None or p["name"] == name)]

    def made(self, suffix: str) -> list[Path]:
        return [self.workspace / f["name"] for f in (self.turn["stats"] or {}).get("files") or []
                if f["name"].lower().endswith(suffix)]


# ------------------------------------------------------------------ checks --


def _pdf_text_and_images(path: Path) -> tuple[list[str], int]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    doc = pdfium.PdfDocument(str(path))
    pages = [doc[i].get_textpage().get_text_range() for i in range(len(doc))]
    images = sum(len(list(doc[i].get_objects(filter=(raw.FPDF_PAGEOBJ_IMAGE,)))) for i in range(len(doc)))
    return pages, images


def check_merge(run: Run) -> dict[str, bool]:
    pdfs = run.made(".pdf")
    out: dict[str, bool] = {"pdf created": bool(pdfs)}
    if not pdfs:
        return out
    pages, images = _pdf_text_and_images(pdfs[-1])
    text = "\n".join(pages)
    out["exp1 content present"] = "Wireless LAN" in text or "WLAN" in text
    out["exp2 content present"] = "Bluetooth" in text
    out["all 10 images"] = images >= 10
    # "1b starts on a new page": the page where exp2's content begins holds no exp1 text.
    first = next((i for i, p in enumerate(pages) if "Bluetooth" in p), None)
    out["exp2 starts a page"] = first is not None and "WLAN" not in pages[first] and "Wireless LAN" not in pages[first]
    out["name on title page"] = bool(pages) and "Pranav Chandak" in pages[0]
    out["no missing images"] = not any("not found" in t["text"] for t in run.tools("create_document"))
    return out


def check_f1(run: Run) -> dict[str, bool]:
    sheets = run.made(".xlsx") + run.made(".csv")
    out: dict[str, bool] = {"spreadsheet created": bool(sheets)}
    if not sheets:
        return out
    rows = _rows(sheets[-1])
    out[">= 20 race rows"] = len(rows) >= 21
    cells = [c.strip().lower() for r in rows[1:] for c in r if isinstance(c, str)]
    out["no team names as drivers"] = not any(c in TEAMS for c in cells)
    blob = " ".join(str(c) for r in rows for c in r) + run.answer
    out["mentions 2025"] = "2025" in blob
    return out


def _rows(path: Path) -> list[list[Any]]:
    if path.suffix.lower() == ".csv":
        return [line.split(",") for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    from openpyxl import load_workbook

    ws = load_workbook(str(path), read_only=True).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def check_research(run: Run) -> dict[str, bool]:
    return {
        "answered": len(run.answer) > 40 and "unable" not in run.answer.lower(),
        "cites a source": bool(re.search(r"https?://\S+", run.answer)),
    }


def check_attachment(run: Run) -> dict[str, bool]:
    return {"names Packet Tracer": "packet tracer" in run.answer.lower()}


def check_xlsx_sum(run: Run) -> dict[str, bool]:
    digits = run.answer.replace(",", "")
    return {"March total 12450.75": "12450.75" in digits}


def check_chat(run: Run) -> dict[str, bool]:
    return {"no tools for a plain question": not run.tools(), "answered": len(run.answer) > 20}


CASES = {c.name: c for c in [
    Case("f1", "Create an Excel file of last season's Formula 1 race results: for every race of the "
               "season, the race name, date, and the drivers who finished 1st, 2nd and 3rd.",
         checks=check_f1),
    Case("research", "Who won the most recent Formula 1 World Drivers' Championship, and by how many "
                     "points? Give the source.", checks=check_research),
    Case("merge", "merge @exp1.pdf and @exp2.pdf into 1 pdf. keep only 1 title page for the merged pdf. "
                  "and in that title page - add name as Pranav Chandak, Class D16 and Roll No. 33. rename "
                  "the experiments as exp 1a and 1b instead of 1 and 2. make sure 1b starts on a new page "
                  "and not continued where 1a ends. keep everything as is. just perform the merge. the text "
                  "and images wherever they are should be intact.",
         uploads=["exp1.pdf", "exp2.pdf"], checks=check_merge),
    Case("attachment", "What apparatus does this experiment require?", uploads=["exp4.pdf"],
         checks=check_attachment),
    Case("xlsx-sum", "What is the total revenue for March in this workbook?", uploads=["sales.xlsx"],
         checks=check_xlsx_sum),
    Case("chat", "In two sentences, what is a VLAN?", checks=check_chat),
]}


def _diagram(path: Path, label: str, n: int) -> Path:
    """A simple network-diagram PNG: boxes joined by lines, labelled."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (640, 360), "white")
    d = ImageDraw.Draw(img)
    hub = (270, 140, 370, 200)
    d.rectangle(hub, outline="black", width=3)
    d.text((282, 160), label[:12], fill="black")
    for i in range(n):
        x = 60 + i * (520 // max(1, n - 1)) if n > 1 else 320
        box = (x - 40, 280, x + 40, 330)
        d.rectangle(box, outline="navy", width=2)
        d.text((x - 30, 298), f"Device {i + 1}", fill="navy")
        d.line((320, 200, x, 280), fill="gray", width=2)
    img.save(path)
    return path


def _lab_pdf(path: Path, number: str, title: str, sections: list[tuple[str, list[str], int]]) -> None:
    """A lab report: a title page alone on page 1, then sections with diagrams."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer

    s = getSampleStyleSheet()
    story: list[Any] = [Spacer(1, 60 * mm), Paragraph(f"Experiment No - {number}", s["Title"]),
                        Paragraph(title, s["Heading1"]), Spacer(1, 20 * mm)]
    story += [Paragraph(f"{k}: ____________", s["Normal"]) for k in ("Name", "Class", "Roll No.", "Sign")]
    story.append(PageBreak())
    img_dir = path.parent / f"_{path.stem}_img"
    img_dir.mkdir(exist_ok=True)
    count = 0
    for heading, paragraphs, images in sections:
        story.append(Paragraph(heading, s["Heading2"]))
        story += [Paragraph(p, s["BodyText"]) for p in paragraphs]
        for _ in range(images):
            count += 1
            png = _diagram(img_dir / f"fig{count}.png", title.split()[0], 2 + count % 3)
            story += [Image(str(png), width=140 * mm, height=79 * mm), Spacer(1, 4 * mm)]
        story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def make_fixtures() -> None:
    """Generate the fixtures once. March in sales.xlsx totals 12450.75."""
    FIXTURES.mkdir(parents=True, exist_ok=True)
    if not (FIXTURES / "exp1.pdf").exists():
        _lab_pdf(FIXTURES / "exp1.pdf", "01", "Wireless LAN Setup", [
            ("Aim", ["Simulating a Wireless LAN setup using Cisco Packet Tracer."], 0),
            ("Procedure", [f"{i}. " + step for i, step in enumerate([
                "Place a wireless router in the workspace.", "Add two laptops and one PC.",
                "Set the SSID to WLAN-Network and enable WPA2-PSK.",
                "Connect each device to WLAN-Network with the passphrase."], 1)], 2),
            ("Verification", ["Ping between the laptops succeeds with an average of 15 ms."], 2),
        ])
    if not (FIXTURES / "exp2.pdf").exists():
        _lab_pdf(FIXTURES / "exp2.pdf", "02", "PAN using Bluetooth", [
            ("Aim", ["To create a Personal Area Network using Bluetooth in Cisco Packet Tracer."], 1),
            ("Steps", [f"{i}. " + step for i, step in enumerate([
                "Select a PC and a Bluetooth speaker.", "Turn on Bluetooth on both devices.",
                "Press Discover on the PC and pair the devices."], 1)], 3),
            ("Result", ["The PC and the speaker are paired and exchange data over Bluetooth."], 2),
        ])
    if not (FIXTURES / "exp4.pdf").exists():
        _lab_pdf(FIXTURES / "exp4.pdf", "04", "Mobile Tower Network", [
            ("Aim and Apparatus", ["Aim: To configure a mobile tower network in 3G and 4G.",
                                   "Apparatus: Cisco Packet Tracer, a computer with 4 GB RAM."], 1),
        ])
    target = FIXTURES / "sales.xlsx"
    if target.exists():
        return
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    data = {"Jan": [1200, 3400.5, 980], "Feb": [2200, 1850.25, 3100],
            "Mar": [4100.5, 2950.25, 5400]}  # March: 12450.75
    for month, values in data.items():
        ws = wb.create_sheet(month)
        ws.append(["region", "revenue"])
        for region, v in zip(("North", "South", "West"), values):
            ws.append([region, v])
    wb.save(str(target))


# ---------------------------------------------------------------- running --


def run_case(client: httpx.Client, case: Case, thinking: bool, data_dir: Path) -> tuple[Run, dict[str, Any]]:
    cid = client.post("/chats", json={"title": f"[bench] {case.name}"}).json()["id"]
    files = []
    for name in case.uploads:
        src = FIXTURES / name
        if not src.is_file():
            raise SystemExit(f"missing fixture {src}")
        files.append(client.post(f"/chats/{cid}/files", files={"file": (name, src.read_bytes())}).json())
    started = time.time()
    o = client.post(f"/chats/{cid}/messages", json={"text": case.prompt, "files": files, "thinking": thinking})
    o.raise_for_status()
    timeline = []
    with client.stream("GET", f"/runs/{o.json()['run']['id']}/events") as r:
        for line in r.iter_lines():
            if not line.startswith("data: "):
                continue
            e = json.loads(line[6:])
            if e["type"] in ("call_start", "tool_start", "tool_end", "model_retry", "resumed",
                             "approval_request", "done", "error"):
                timeline.append((round(time.time() - started, 1), e["type"], e.get("name", e.get("message") or "")))
            if e["type"] in ("done", "error"):
                break
    turn = client.get(f"/chats/{cid}").json()["turns"][-1]
    run = Run(case.name, thinking, cid, turn, data_dir / "chats" / cid)
    checks = case.checks(run) if case.checks else {}
    stats = turn["stats"] or {}
    summary = {
        "case": case.name, "thinking": thinking, "chat": cid, "status": turn["status"],
        "seconds": round(stats.get("duration_ms", 0) / 1000, 1), "model_calls": stats.get("model_calls"),
        "context": stats.get("context_used"), "tools": [t["name"] for t in run.tools()],
        "failed_tools": sum(1 for t in run.tools() if t["status"] == "error"),
        "checks": checks, "passed": all(checks.values()) if checks else None,
        "answer": run.answer[:600], "timeline": timeline,
    }
    return run, summary


def trace(run: Run, summary: dict[str, Any]) -> str:
    lines = [f"# {run.case} (thinking {'on' if run.thinking else 'off'})", "",
             "```json", json.dumps({k: v for k, v in summary.items() if k != "timeline"}, indent=1), "```",
             "", "## Timeline", ""] + [f"- {a}s {b} {c}" for a, b, c in summary["timeline"]] + ["", "## Steps", ""]
    for i, p in enumerate(run.parts, 1):
        if p["kind"] == "reasoning":
            lines += [f"### {i}. reasoning ({len(p['text'])} chars)", "```", p["text"], "```", ""]
        elif p["kind"] == "tool":
            lines += [f"### {i}. tool `{p['name']}` -> {p['status']}", "```json",
                      json.dumps(p["args"], indent=1, ensure_ascii=False)[:8000], "```",
                      f"result ({len(p['text'])} chars):", "```", p["text"][:4000], "```", ""]
        else:
            lines += [f"### {i}. {p['kind']}", p.get("text") or json.dumps(p)[:600], ""]
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("cases", nargs="*", default=list(CASES), help=f"of {list(CASES)}")
    parser.add_argument("--thinking", choices=["off", "on", "both"], default="both")
    parser.add_argument("--url", default="http://127.0.0.1:8000")
    args = parser.parse_args()

    make_fixtures()
    client = httpx.Client(base_url=f"{args.url}/api", timeout=1800)
    settings = client.get("/settings").json()
    data_dir = Path(settings["paths"]["data"])
    out = data_dir / "bench" / time.strftime("%Y%m%d-%H%M%S")
    out.mkdir(parents=True, exist_ok=True)
    modes = {"off": [False], "on": [True], "both": [False, True]}[args.thinking]

    results = []
    for name in args.cases:
        for thinking in modes:
            run, summary = run_case(client, CASES[name], thinking, data_dir)
            (out / f"{name}-think{int(thinking)}.md").write_text(trace(run, summary), encoding="utf-8")
            results.append(summary)
            mark = {True: "PASS", False: "FAIL", None: "----"}[summary["passed"]]
            failed = [k for k, v in summary["checks"].items() if not v]
            print(f"{mark} {name:<11} thinking={'on ' if thinking else 'off'} {summary['seconds']:>6}s "
                  f"calls={summary['model_calls']} tools={len(summary['tools'])} "
                  f"failed_tools={summary['failed_tools']}" + (f"  missing: {failed}" if failed else ""),
                  flush=True)
    (out / "summary.json").write_text(json.dumps({"settings": settings["values"], "results": results},
                                                 indent=1, default=str), encoding="utf-8")
    passed = sum(1 for r in results if r["passed"])
    print(f"\n{passed}/{len(results)} passed. Traces: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
