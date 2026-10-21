"""The HTTP API. Chats, turns, runs, settings -- all under `/api`.

One active run at a time, a run that outlives the frontend, and a chat tree
whose model history lives in Strands snapshots.

    send    a turn under the selected leaf, restoring the parent's after_snap
    retry   last turn only: a sibling with the same parent and restore point,
            sampled wider and told what the earlier versions said
    edit    any turn: delete it, its versions and everything below, then run
            the new text from the same restore point
    <  />   select another version at a position; nothing is recomputed
    cancel  agent.cancel(); the turn keeps its partial transcript but no
            snapshot, so it can never become the parent of a follow-up

Everything is under `/api` so the frontend can own `/settings`, `/skills` and
`/mcp` as page paths without the dev proxy forwarding them to the backend.
"""

from __future__ import annotations

import asyncio
import json
import logging
import mimetypes
import os
import re
import shutil
import sys
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any

from fastapi import APIRouter, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel, Field

from agent.core import form_fill, llama, paths
from agent.core.agent_factory import RetryPlan, build_agent, clock_note, retry_hint
from agent.core.chats import (
    CANCELLED,
    COMPLETED,
    FAILED,
    RUNNING,
    ChatStore,
    TreeError,
)
from agent.core.mcp_manager import McpManager
from agent.core.memory_manager import MemoryManager
from agent.core.session_manager import SessionManager
from agent.core.settings import (
    TIERS,
    Settings,
    SettingsError,
    changed_restart_keys,
    flatten,
    load_settings,
    parse_settings,
    restart_tier_values,
    set_path,
    write_settings,
)
from agent.core.page_tools import import_pdf, page_note, page_tools, write_table
from agent.core.shell_env import model_path
from agent.core.skill_manager import SkillManager, reader_command
from agent.core.tool_manager import ToolManager
from server.security import EXTENSION_SCHEMES, Pairing, guard
from server.events import _aiter, translate
from server.runs import Run, RunConflict, RunManager
from server.transcript import answer_of

logger = logging.getLogger(__name__)


# ---------------------------------------------------------- request bodies --


class CreateChat(BaseModel):
    title: str = ""
    #: `web`, or `extension` for a chat started in the browser side panel.
    source: str = "web"
    page_url: str = ""
    page_title: str = ""


class Page(BaseModel):
    """The tab a side-panel message is about. Its presence turns on the page tools."""

    url: str = ""
    title: str = ""


class RenameChat(BaseModel):
    title: str


class SendMessage(BaseModel):
    text: str = ""
    files: list[dict[str, Any]] = Field(default_factory=list)
    #: Passages selected in the preview panel: {"file", "path", "quote"}.
    refs: list[dict[str, Any]] = Field(default_factory=list)
    thinking: bool | None = None
    page: Page | None = None


class RetryTurn(BaseModel):
    thinking: bool | None = None
    page: Page | None = None


class EditTurn(BaseModel):
    text: str
    files: list[dict[str, Any]] | None = None
    refs: list[dict[str, Any]] | None = None
    thinking: bool | None = None
    page: Page | None = None


class SelectTurn(BaseModel):
    delta: int = 0


class Answer(BaseModel):
    allow: bool


class ProfileBody(BaseModel):
    text: str


class FillBody(BaseModel):
    #: [{id, label, type, options?, required?}]; never the fields' contents.
    fields: list[dict[str, Any]]
    title: str = ""


class ImportFile(BaseModel):
    url: str


class TableBody(BaseModel):
    rows: list[list[str]]
    title: str = ""
    url: str = ""
    format: str = "xlsx"
    #: Where to save; the extension's setting. Default: Downloads/Atlas.
    folder: str = ""


class CallResult(BaseModel):
    """What the client did for a client call: anything, or `error`."""

    model_config = {"extra": "allow"}


# ------------------------------------------------------------------ state --


class Atlas:
    """The backend's state, held outside the app for testability."""

    def __init__(self, settings_path: Path) -> None:
        self.settings_path = Path(settings_path)
        #: As the process started. Restart-tier keys keep these values for the
        #: life of the process even after settings.json changes, because the
        #: running llama-server was built with them.
        self.boot = load_settings(self.settings_path)
        self.settings = self.boot
        self.store = ChatStore(paths.database(self.settings))
        self.runs = RunManager()
        #: MCP connections live as long as the process, not as long as a turn.
        self.mcp = McpManager(self.settings)
        self._modalities: dict[str, bool] | None = None
        if n := self.store.mark_interrupted():
            logger.info("%d turn(s) left running by a previous process are now interrupted", n)
        logger.info("backend ready (root=%s, db=%s)", paths.ROOT, self.store.db_path)

    def modalities(self) -> dict[str, bool]:
        """What the loaded model can take in (vision, audio), from llama-server.

        Asked once and remembered: it changes only with the model, which needs
        a restart, and a restart starts a new backend. Not remembered while
        llama-server is still loading.
        """
        if self._modalities is None:
            try:
                import httpx

                props = httpx.get(f"http://127.0.0.1:{self.boot.ports.llama}/props", timeout=2,
                                  headers=llama.auth_headers()).json()
                self._modalities = {k: bool(v) for k, v in (props.get("modalities") or {}).items()}
            except Exception:  # noqa: BLE001 - not up yet: ask again next time
                return {}
        return self._modalities

    def close(self) -> None:
        self.mcp.stop_all()
        self.store.close()

    # ------------------------------------------------------------ helpers --

    def require_chat(self, chat_id: str) -> dict[str, Any]:
        chat = self.store.get_chat(chat_id)
        if chat is None:
            raise HTTPException(404, f"no chat {chat_id}")
        return chat

    def require_turn(self, turn_id: str) -> dict[str, Any]:
        turn = self.store.get_turn(turn_id)
        if turn is None:
            raise HTTPException(404, f"no turn {turn_id}")
        return turn

    def chat_dir(self, chat_id: str) -> Path:
        path = paths.chat_dir(self.settings, chat_id)
        path.mkdir(parents=True, exist_ok=True)
        return path

    def ensure_run_slot(self) -> None:
        try:
            self.runs.ensure_available()
        except RunConflict as exc:
            raise HTTPException(409, str(exc)) from exc

    def thinking(self, override: bool | None) -> bool:
        return self.settings.sampling.thinking if override is None else override

    def view(self, chat_id: str) -> dict[str, Any]:
        """The visible path with everything the UI needs to draw it. No snapshots read."""
        chat = self.require_chat(chat_id)
        run = self.runs.active_for_chat(chat_id)
        turns = []
        for turn in self.store.path(chat_id):
            versions = [v["id"] for v in self.store.siblings_of(turn["id"])]
            turns.append(_turn_view(turn, versions))
        return {**chat, "turns": turns, "active_run": _run_view(run) if run else None}

    # -------------------------------------------------------------- runs --

    def launch(
        self, turn: dict[str, Any], *, thinking: bool, retry: RetryPlan | None = None,
        page: Page | None = None,
    ) -> Run:
        try:
            run = self.runs.start(turn["chat_id"], turn["id"],
                                  context_size=self.settings.model.context_size)
        except RunConflict as exc:
            raise HTTPException(409, str(exc)) from exc
        run.page = page.model_dump() if page else None
        run.task = asyncio.create_task(self.execute(run, turn, thinking=thinking, retry=retry))
        return run

    def _page_tools(self, run: Run) -> list[Any]:
        """The page tools, for a turn sent from the side panel with its tab attached."""
        if not run.page:
            return []

        def read() -> None:
            run.tainted = True

        workspace = self.chat_dir(run.chat_id)
        return page_tools(run.client_call, workspace, lambda p: model_path(str(p)), on_read=read)

    async def execute(
        self, run: Run, turn: dict[str, Any], *, thinking: bool, retry: RetryPlan | None,
    ) -> None:
        """The run body. One task per run; the frontend is not involved."""
        settings = self.settings
        files = json.loads(turn["files_json"] or "[]")
        # A passage selected in the preview names its file relative to the
        # workspace; the model is given the full path, ready for edit_document.
        workspace = self.chat_dir(run.chat_id)
        refs = [{**r, "path": r.get("path") or str(workspace / str(r.get("file", "")))}
                for r in json.loads(turn.get("refs_json") or "[]") if isinstance(r, dict)]
        media = [f for f in files if _media_kind(f) and Path(str(f.get("path", ""))).is_file()]
        text = _agent_text(turn["user_text"], [f for f in files if f not in media], refs,
                           reader=lambda p, shown: reader_command(settings, p, shown),
                           workspace=workspace)
        if guidance := _mention_guidance(turn["user_text"], settings, self.mcp, workspace):
            text = f"{text}\n\n{guidance}".strip()
        if links := _link_guidance(turn["user_text"], settings):
            text = f"{text}\n\n{links}".strip()
        if run.page:
            text = f"{text}\n\n{page_note(run.page)}".strip()
        # Only the model sees this; the chat shows the user's own text.
        text = f"{text}\n\n{clock_note()}".strip()
        hint = retry_hint(settings, retry) if retry else ""
        sm = SessionManager(settings, run.chat_id)
        # Recorded per turn, so each version shows the model that wrote it.
        run.transcript.stats["model"] = settings.model.name
        #: What the workspace held before the run, so the files this turn
        #: created or changed can be offered as downloads afterwards.
        run.workspace_before = _scan(self.chat_dir(run.chat_id))
        try:
            manager = sm.manager()
            # Building restores a snapshot from disk: blocking work that would
            # stall every open SSE stream if it ran on the event loop.
            agent = await asyncio.to_thread(
                build_agent, settings, chat_id=run.chat_id, thinking=thinking,
                restore_snapshot_id=turn["before_snap"], session_manager=manager,
                retry=retry, observer=run.observe, tools=self.mcp.agent_tools() + self._page_tools(run),
                approver=lambda kind, detail: run.request_approval(
                    kind, detail, settings.approvals.timeout),
                sensitive=self.mcp.is_sensitive,
                tainted=lambda: run.tainted,
            )
            run.agent = agent
            if run.cancel.is_set():
                raise asyncio.CancelledError
            run.emit({"type": "run_started", "chat_id": run.chat_id, "turn_id": run.turn_id})

            start = len(agent.messages)
            prompt_text = f"{text}\n\n{hint}" if hint else text
            # Images and recordings go to the model as content, after the text
            # and in separate blocks (a combined block silently drops the text).
            prompt: Any = prompt_text
            if media:
                blocks = await asyncio.to_thread(_media_blocks, media)
                prompt = ([{"text": prompt_text}] if prompt_text else []) + blocks
            async for out in translate(_aiter(agent.stream_async(prompt))):
                if run.cancel.is_set():
                    raise asyncio.CancelledError
                run.emit(out)

            if hint:
                _strip_hint(agent.messages, start, hint)
            snap = await manager.save_snapshot(agent, is_latest=False)
            self._finish(run, COMPLETED, after_snap=snap)

        except (asyncio.CancelledError, KeyboardInterrupt):
            # No snapshot: a cancelled turn can be retried or edited, never
            # continued, so a half-finished tool call is never built upon.
            self._finish(run, CANCELLED)

        except Exception as exc:  # noqa: BLE001
            logger.exception("run %s failed", run.id)
            self._finish(run, FAILED, error=f"{type(exc).__name__}: {exc}")

        finally:
            run.agent = None
            self.runs.adopt(run)

    def _finish(self, run: Run, status: str, *, after_snap: str | None = None,
                error: str | None = None) -> None:
        stats = run.transcript.final_stats()
        try:
            stats["files"] = _changed(run.workspace_before,
                                      _scan(self.chat_dir(run.chat_id)))
        except OSError:
            logger.warning("could not scan the workspace of chat %s", run.chat_id, exc_info=True)
        try:
            self.store.finish_turn(run.turn_id, status, after_snap=after_snap,
                                   render=run.transcript.parts, stats=stats, error=error)
            self.store.touch(run.chat_id)
            title = (self.store.get_chat(run.chat_id) or {}).get("title") or "untitled"
            logger.info("run %s %s · chat “%s” · %s model calls · %.1fs", run.id[:4], status,
                        title, stats.get("model_calls", 0), (stats.get("duration_ms") or 0) / 1000)
        except Exception:  # noqa: BLE001 - the turn may have been cut by an edit
            logger.exception("could not record the end of turn %s", run.turn_id)
        run.finish(status, error)


# -------------------------------------------------------------------- app --


def create_app(settings_path: str | Path | None = None) -> FastAPI:
    state = Atlas(paths.resolve(settings_path or "settings.json"))

    @asynccontextmanager
    async def lifespan(_: FastAPI):
        loop = asyncio.get_running_loop()
        if sys.platform == "win32" and "Proactor" not in type(loop).__name__:
            # The Selector loop cannot start subprocesses (see server/main.py).
            # `uvicorn --reload` is the usual cause.
            logger.error("running on %s: shell tools will not work. Start the backend "
                         "with `atlas start`, or uvicorn WITHOUT --reload.",
                         type(loop).__name__)
        # Connected in the background: a slow `npx` download must not hold up
        # the API, and a chat started meanwhile simply has no MCP tools yet.
        loop.run_in_executor(None, state.mcp.start_all)
        yield
        for run in list(state.runs._runs.values()):  # noqa: SLF001
            if run.task and not run.task.done():
                state.runs.cancel(run.id)
        state.close()

    app = FastAPI(title="Atlas", lifespan=lifespan)
    app.state.atlas = state
    pairing = Pairing(paths.data_dir(state.settings) / "pairing-token")
    ports = state.boot.ports
    trusted = {f"http://{h}:{p}" for h in ("127.0.0.1", "localhost") for p in (ports.backend, ports.client)}
    app.middleware("http")(guard(pairing, trusted))
    api = APIRouter(prefix="/api")

    # ------------------------------------------------------------ pairing --

    @api.get("/pairing")
    async def get_pairing(request: Request) -> dict[str, Any]:
        """The token to paste into the browser extension. Not for the extension itself."""
        if request.headers.get("origin", "").startswith(EXTENSION_SCHEMES):
            raise HTTPException(403, "the extension cannot read the pairing token")
        return {"token": pairing.token}

    @api.post("/pairing/reset")
    async def reset_pairing(request: Request) -> dict[str, Any]:
        """A new token: the extension has to be paired again."""
        if request.headers.get("origin", "").startswith(EXTENSION_SCHEMES):
            raise HTTPException(403, "the extension cannot reset the pairing token")
        return {"token": pairing.reset()}

    @api.get("/profile")
    async def get_profile(request: Request) -> dict[str, Any]:
        """The user's details for fill suggestions. The web UI's, never the extension's."""
        if request.headers.get("origin", "").startswith(EXTENSION_SCHEMES):
            raise HTTPException(403, "the extension cannot read your details")
        return {"text": form_fill.read_profile(state.settings)}

    @api.put("/profile")
    async def put_profile(body: ProfileBody, request: Request) -> dict[str, Any]:
        if request.headers.get("origin", "").startswith(EXTENSION_SCHEMES):
            raise HTTPException(403, "edit your details in Atlas, not in the extension")
        await asyncio.to_thread(form_fill.write_profile, state.settings, body.text)
        return {"text": body.text}

    @api.post("/extension/fill")
    async def fill(body: FillBody) -> dict[str, Any]:
        """Suggested values for a form's fields, from the user's details. Nothing is filled here."""
        try:
            return await asyncio.to_thread(form_fill.suggest, state.settings, body.fields, body.title)
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc

    @api.post("/extension/table")
    async def save_table(body: TableBody) -> dict[str, Any]:
        """A table from a web page, saved as xlsx or csv. Code only: no model involved."""
        try:
            path = await asyncio.to_thread(write_table, body.rows, body.title, body.url, body.format,
                                           body.folder)
        except (OSError, ValueError) as exc:
            raise HTTPException(422, str(exc)) from exc
        return {"path": str(path), "rows": len(body.rows)}

    # ------------------------------------------------------------- health --

    @api.get("/health")
    async def health() -> dict[str, Any]:
        run = state.runs.active()
        if run is None:
            return {"status": "ok", "active_run": None}
        chat = state.store.get_chat(run.chat_id) or {}
        return {"status": "ok", "active_run": {**_run_view(run), "chat_title": chat.get("title") or ""}}

    # -------------------------------------------------------------- chats --

    @api.get("/chats")
    async def list_chats() -> list[dict[str, Any]]:
        return state.store.list_chats()

    @api.post("/chats")
    async def create_chat(body: CreateChat | None = None) -> dict[str, Any]:
        if body is None:
            return state.store.create_chat("")
        source = body.source if body.source in ("web", "extension") else "web"
        return state.store.create_chat(body.title or "", source=source, page_url=body.page_url[:2000],
                                       page_title=body.page_title[:300])

    @api.get("/chats/{chat_id}")
    async def get_chat(chat_id: str) -> dict[str, Any]:
        return state.view(chat_id)

    @api.patch("/chats/{chat_id}")
    async def rename_chat(chat_id: str, body: RenameChat) -> dict[str, Any]:
        state.require_chat(chat_id)
        state.store.set_title(chat_id, body.title.strip()[:200])
        return state.store.get_chat(chat_id)  # type: ignore[return-value]

    @api.delete("/chats/{chat_id}")
    async def delete_chat(chat_id: str) -> dict[str, Any]:
        state.require_chat(chat_id)
        if state.runs.active_for_chat(chat_id):
            raise HTTPException(409, "cannot delete a chat with an active run")
        state.store.delete_chat(chat_id)
        await SessionManager(state.settings, chat_id).delete_all()
        shutil.rmtree(paths.chat_dir(state.settings, chat_id), ignore_errors=True)
        return {"deleted": chat_id}

    @api.get("/chats/{chat_id}/export")
    async def export_chat(chat_id: str) -> dict[str, Any]:
        """Everything needed to restore the chat elsewhere, version tree and
        Strands snapshots included, so an imported chat can be continued."""
        state.require_chat(chat_id)
        rows = state.store.export_rows(chat_id)
        snaps = await SessionManager(state.settings, chat_id).read_snapshots(
            t["after_snap"] for t in rows["turns"])
        return {"format": "atlas-chat", "version": 1, **rows, "snapshots": snaps}

    @api.post("/chats/import")
    async def import_chat(body: dict[str, Any]) -> dict[str, Any]:
        if body.get("format") != "atlas-chat" or not isinstance(body.get("turns"), list):
            raise HTTPException(422, "not an Atlas chat export")
        chat_id = state.store.import_rows(body.get("chat") or {}, body["turns"])
        await SessionManager(state.settings, chat_id).write_snapshots(body.get("snapshots") or {})
        return state.store.get_chat(chat_id)  # type: ignore[return-value]

    # --------------------------------------------------------------- files --

    @api.post("/chats/{chat_id}/files")
    async def upload(chat_id: str, request: Request) -> dict[str, Any]:
        state.require_chat(chat_id)
        form = await request.form()
        upload_file = form.get("file")
        if upload_file is None or isinstance(upload_file, str):
            raise HTTPException(422, "no file part")
        name = Path(upload_file.filename or "upload.bin").name
        target = state.chat_dir(chat_id) / name

        # Nothing is parsed at upload: the file is saved and the model reads it
        # on demand through a skill. Copying is still blocking I/O, so it runs
        # off the event loop and cannot freeze an active run's stream.
        def _save() -> None:
            with target.open("wb") as fh:
                shutil.copyfileobj(upload_file.file, fh)

        await asyncio.to_thread(_save)
        return {"name": name, "path": str(target), "size": target.stat().st_size,
                "type": mimetypes.guess_type(name)[0] or "application/octet-stream"}

    @api.post("/chats/{chat_id}/files/import")
    async def import_file(chat_id: str, body: ImportFile) -> dict[str, Any]:
        """A PDF the extension's tab shows, as an attachment.

        The browser's PDF viewer runs no content script, so the page tools
        cannot read it; the file goes to the workspace for the pdf skill. A
        `file://` PDF is copied from disk (so no "allow access to file URLs"
        switch is needed); a web one is downloaded, for when the extension could
        not fetch it itself.
        """
        state.require_chat(chat_id)
        try:
            target = await asyncio.to_thread(import_pdf, body.url, state.chat_dir(chat_id))
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
        return {"name": target.name, "path": str(target), "size": target.stat().st_size,
                "type": "application/pdf"}

    def workspace_file(chat_id: str, name: str) -> Path:
        state.require_chat(chat_id)
        root = paths.chat_dir(state.settings, chat_id).resolve()
        target = (root / name).resolve()
        if root not in target.parents:
            raise HTTPException(400, "path escapes the chat directory")
        if not target.is_file():
            raise HTTPException(404, "no such file")
        return target

    @api.delete("/chats/{chat_id}/files/{name:path}")
    async def remove_upload(chat_id: str, name: str) -> dict[str, Any]:
        """Remove an upload the user took off a message before sending it.

        Refused for a file any sent message used: that message's history, and
        any document made from it, still point at it.
        """
        target = workspace_file(chat_id, name)
        for row in state.store.export_rows(chat_id)["turns"]:
            if any(f.get("name") == target.name for f in json.loads(row["files_json"] or "[]")):
                raise HTTPException(409, f"{target.name} was sent with a message and is kept")
        target.unlink()
        return {"deleted": name}

    @api.get("/chats/{chat_id}/workspace")
    async def workspace(chat_id: str) -> list[dict[str, Any]]:
        """Every file in the chat's workspace, newest first. Outlives any edit."""
        state.require_chat(chat_id)
        root = state.chat_dir(chat_id)
        entries = [{"name": name, "size": size, "path": str(root / name),
                    "modified": mtime // 1_000_000,
                    "type": mimetypes.guess_type(name)[0] or "application/octet-stream"}
                   for name, (size, mtime) in _scan(root).items()]
        return sorted(entries, key=lambda e: e["modified"], reverse=True)

    @api.get("/chats/{chat_id}/files/{name:path}")
    async def download(chat_id: str, name: str, download: bool = False) -> FileResponse:
        """Served inline for the preview panel; `?download=1` saves it instead."""
        target = workspace_file(chat_id, name)
        if download:
            return FileResponse(target, filename=target.name)
        return FileResponse(target, headers={"Cache-Control": "no-cache"})

    @api.get("/chats/{chat_id}/preview/{name:path}")
    async def preview(chat_id: str, name: str) -> dict[str, Any]:
        """A spreadsheet as rows, for the preview panel. Capped, and says so."""
        target = workspace_file(chat_id, name)
        if target.suffix.lower() not in (".xlsx", ".xlsm"):
            raise HTTPException(415, "only spreadsheets have a server-side preview")
        return await asyncio.to_thread(_sheet_preview, target)

    # -------------------------------------------------------------- turns --

    @api.post("/chats/{chat_id}/messages")
    async def send(chat_id: str, body: SendMessage) -> dict[str, Any]:
        state.require_chat(chat_id)
        if not body.text.strip() and not body.files:
            raise HTTPException(422, "empty message")
        # Checked before inserting, so a 409 never leaves a stranded turn.
        state.ensure_run_slot()

        leaf = state.store.leaf(chat_id)
        if leaf is not None and leaf["status"] != COMPLETED:
            raise HTTPException(409, "the last answer did not complete: retry or edit it first")
        turn = state.store.insert_turn(
            chat_id, body.text, parent_turn_id=leaf["id"] if leaf else None,
            before_snap=leaf["after_snap"] if leaf else None, files=body.files, refs=body.refs,
        )
        if first_words := " ".join(body.text.split())[:60]:
            state.store.set_title(chat_id, first_words, only_if_empty=True)
        run = state.launch(turn, thinking=state.thinking(body.thinking), page=body.page)
        return {"run": _run_view(run), "chat": state.view(chat_id)}

    @api.post("/turns/{turn_id}/retry")
    async def retry(turn_id: str, body: RetryTurn | None = None) -> dict[str, Any]:
        turn = state.require_turn(turn_id)
        if turn["status"] == RUNNING:
            raise HTTPException(409, "that turn is still running")
        if not state.store.is_leaf(turn_id):
            raise HTTPException(409, "only the last answer can be retried; edit earlier turns")
        state.ensure_run_slot()

        # Newest first: the version the user just rejected is the one most
        # worth steering away from.
        previous = [answer_of(json.loads(v["render_json"] or "[]"))
                    for v in reversed(state.store.siblings_of(turn_id))]
        new = state.store.insert_turn(
            turn["chat_id"], turn["user_text"], parent_turn_id=turn["parent_turn_id"],
            before_snap=turn["before_snap"], files=json.loads(turn["files_json"] or "[]"),
            refs=json.loads(turn.get("refs_json") or "[]"),
        )
        run = state.launch(new, thinking=state.thinking(body.thinking if body else None),
                           retry=RetryPlan(previous=[p for p in previous if p.strip()]),
                           page=body.page if body else None)
        return {"run": _run_view(run), "chat": state.view(turn["chat_id"])}

    @api.post("/turns/{turn_id}/edit")
    async def edit(turn_id: str, body: EditTurn) -> dict[str, Any]:
        turn = state.require_turn(turn_id)
        files = body.files if body.files is not None else json.loads(turn["files_json"] or "[]")
        refs = body.refs if body.refs is not None else json.loads(turn.get("refs_json") or "[]")
        if not body.text.strip() and not files:
            raise HTTPException(422, "empty message")
        state.ensure_run_slot()

        # Destructive by design: the old future answered a question that is
        # no longer being asked.
        try:
            old, snaps = state.store.cut_from(turn_id)
        except TreeError as exc:
            raise HTTPException(404, str(exc)) from exc
        await SessionManager(state.settings, old["chat_id"]).delete_snapshots(snaps)
        # Files in the workspace are never touched by an edit: documents the
        # agent made or changed for the deleted turns stay, and stay listed.
        new = state.store.insert_turn(
            old["chat_id"], body.text, parent_turn_id=old["parent_turn_id"],
            before_snap=old["before_snap"], files=files, refs=refs,
        )
        run = state.launch(new, thinking=state.thinking(body.thinking), page=body.page)
        return {"run": _run_view(run), "chat": state.view(old["chat_id"])}

    @api.post("/turns/{turn_id}/select")
    async def select(turn_id: str, body: SelectTurn) -> dict[str, Any]:
        """`<` / `>`: pick another version. One query, no snapshot read."""
        turn = state.require_turn(turn_id)
        if state.runs.active_for_chat(turn["chat_id"]):
            raise HTTPException(409, "wait for the current answer to finish")
        versions = state.store.siblings_of(turn_id)
        ids = [v["id"] for v in versions]
        index = max(0, min(ids.index(turn_id) + body.delta, len(ids) - 1))
        state.store.select(ids[index])
        return state.view(turn["chat_id"])

    # --------------------------------------------------------------- runs --

    @api.get("/runs/{run_id}")
    async def get_run(run_id: str) -> dict[str, Any]:
        run = state.runs.get(run_id)
        if run is None:
            raise HTTPException(404, f"no run {run_id}")
        return {**_run_view(run), "error": run.error, "events": len(run.events)}

    @api.get("/runs/{run_id}/events")
    async def run_events(run_id: str, after: int = 0) -> StreamingResponse:
        run = state.runs.get(run_id)
        if run is None:
            raise HTTPException(404, f"no run {run_id}")

        async def sse() -> Any:
            # Replay then follow, so a reloaded client resumes with no gaps and
            # no duplicates. `after` is the first seq the client has NOT seen.
            async for event in state.runs.stream(run, after):
                yield f"data: {json.dumps(event)}\n\n"

        return StreamingResponse(sse(), media_type="text/event-stream",
                                 headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})

    @api.post("/runs/{run_id}/approvals/{approval_id}")
    async def answer_approval(run_id: str, approval_id: str, body: Answer) -> dict[str, Any]:
        run = state.runs.get(run_id)
        if run is None:
            raise HTTPException(404, f"no run {run_id}")
        if not run.answer(approval_id, body.allow):
            raise HTTPException(409, "that question was already answered or has expired")
        return {"id": approval_id, "decision": "allowed" if body.allow else "denied"}

    @api.post("/runs/{run_id}/calls/{call_id}")
    async def answer_call(run_id: str, call_id: str, body: CallResult) -> dict[str, Any]:
        """The side panel's result for a page operation the run asked for."""
        run = state.runs.get(run_id)
        if run is None:
            raise HTTPException(404, f"no run {run_id}")
        if not run.resolve_call(call_id, body.model_dump()):
            raise HTTPException(409, "that call was already answered or has timed out")
        return {"id": call_id}

    @api.post("/runs/{run_id}/cancel")
    async def cancel_run(run_id: str) -> dict[str, Any]:
        run = state.runs.get(run_id)
        if run is None:
            raise HTTPException(404, f"no run {run_id}")
        state.runs.cancel(run_id)
        return _run_view(run)

    # ------------------------------------------------------------ settings --

    def settings_view() -> dict[str, Any]:
        on_disk = json.loads(state.settings_path.read_text(encoding="utf-8"))
        return {
            "values": flatten(parse_settings(on_disk).model_dump(mode="json")),
            "tiers": {k: t.value for k, t in TIERS.items()},
            "pending_restart": changed_restart_keys(on_disk, _launched_restart_values(state)),
            "tools": ToolManager(state.settings).describe(),
            "memory": MemoryManager(state.settings).describe(),
            "skills": SkillManager(state.settings).catalogue(),
            "mcp": state.mcp.describe(),
            "paths": {"root": str(paths.ROOT), "data": str(paths.data_dir(state.settings))},
            "modalities": state.modalities(),
        }

    @api.get("/settings")
    async def get_settings() -> dict[str, Any]:
        return settings_view()

    @api.put("/settings")
    async def put_settings(patch: dict[str, Any]) -> dict[str, Any]:
        """Write any key. LIVE keys apply on the next turn; RESTART keys are
        saved to settings.json and take effect at the next `atlas start`."""
        unknown = sorted(k for k in patch if k not in TIERS)
        if unknown:
            raise HTTPException(422, f"unknown setting(s): {unknown}")
        on_disk = json.loads(state.settings_path.read_text(encoding="utf-8"))
        for dotted, value in patch.items():
            set_path(on_disk, dotted, value)
        try:
            write_settings(state.settings_path, on_disk)
        except SettingsError as exc:
            raise HTTPException(422, str(exc)) from exc
        state.settings = _live_settings(on_disk, state.boot)
        # Servers whose arguments hold {atlas:folders} reconnect with the new list.
        await asyncio.to_thread(state.mcp.set_settings, state.settings)
        return settings_view()

    # ------------------------------------------------------------- models --

    @api.get("/models")
    async def list_models() -> dict[str, Any]:
        """The models in the models folder, for the Settings picker."""
        found = await asyncio.to_thread(llama.discover, paths.models_dir(state.settings))
        return {"models": [m.describe() for m in found], "loaded": state.boot.model.name,
                "selected": state.settings.model.name}

    # ------------------------------------------------------------ restart --

    @api.get("/restart")
    async def restart_status() -> dict[str, Any]:
        """How the last restart went: restarting, ready or failed (with why)."""
        status = os.environ.get("ATLAS_RESTART_STATUS")
        available = bool(os.environ.get("ATLAS_RESTART_REQUEST"))
        try:
            data = json.loads(Path(status).read_text(encoding="utf-8")) if status else {}
        except (OSError, json.JSONDecodeError):
            data = {}
        return {"available": available, **data}

    @api.post("/restart")
    async def restart() -> dict[str, Any]:
        """Ask `atlas start` to restart llama-server and this backend.

        The CLI watches for the request file; this process is replaced, so the
        UI polls GET /restart until it answers again.
        """
        request = os.environ.get("ATLAS_RESTART_REQUEST")
        if not request:
            raise HTTPException(409, "Atlas was not started with `atlas start`, so it cannot restart itself. "
                                     "Stop it and start it again by hand.")
        if state.runs.active():
            raise HTTPException(409, "An answer is being written. Stop it or wait, then restart.")
        Path(request).write_text(str(int(time.time() * 1000)), encoding="utf-8")
        return {"ok": True}

    app.state.settings_view = settings_view
    app.include_router(api)

    # Skills and MCP routes, also under /api.
    from server.extensions import register as register_extensions  # noqa: PLC0415

    register_extensions(app, state)
    _serve_client(app)
    return app


def _serve_client(app: FastAPI) -> None:
    """The built web UI (`client/dist`), on the API's own port and origin.

    `atlas start` uses this instead of the Vite dev server, so normal use needs
    no Node. Any path that is not a file in the build gets `index.html`, so a
    reload on `/settings/agent` still opens the app. Registered after the API,
    whose routes win. Looked up per request: `atlas start` may still be
    building it when the backend comes up. With `--dev` (ATLAS_DEV_UI set),
    pages redirect to Vite, so the backend's port never shows a stale build.
    """
    dist = paths.ROOT / "client" / "dist"
    dev_ui = os.environ.get("ATLAS_DEV_UI")

    @app.get("/{path:path}", include_in_schema=False)
    async def client(path: str, request: Request) -> Any:
        if path.startswith("api/"):
            raise HTTPException(404, f"no route /{path}")
        if dev_ui:  # `atlas start --dev`: the live UI is Vite's, on its own port
            query = f"?{request.url.query}" if request.url.query else ""
            return RedirectResponse(f"{dev_ui}/{path}{query}")
        if not (dist / "index.html").is_file():
            raise HTTPException(503, "The web UI is not built: run `npm run build` in client/, or `atlas start --dev`.")
        target = (dist / path).resolve()
        if path and target.is_file() and dist.resolve() in target.parents:
            return FileResponse(target)
        return FileResponse(dist / "index.html")


# ---------------------------------------------------------------- helpers --


def _live_settings(on_disk: dict[str, Any], boot: Settings) -> Settings:
    """The file's LIVE values, with RESTART values held at what the process started with."""
    merged = json.loads(json.dumps(on_disk))
    for key, value in restart_tier_values(boot.model_dump(mode="json")).items():
        set_path(merged, key, value)
    return parse_settings(merged)


def _launched_restart_values(state: Atlas) -> dict[str, Any]:
    """What the running servers were started with.

    launched.json when the CLI wrote one (its `restart_settings` block is flat,
    `{"model.path": ...}`), otherwise this process's own start-up values.
    """
    launched = Path(os.environ.get("ATLAS_LAUNCHED") or paths.ROOT / "launched.json")
    if launched.is_file():
        try:
            flat = json.loads(launched.read_text(encoding="utf-8")).get("restart_settings")
            if isinstance(flat, dict) and flat:
                return flat
        except (OSError, json.JSONDecodeError):
            pass
    return restart_tier_values(state.boot.model_dump(mode="json"))


_IMAGE = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg", "gif": "gif", "webp": "webp"}
_AUDIO = {"wav": "wav", "mp3": "mp3"}


def _media_kind(f: dict[str, Any]) -> str | None:
    ext = Path(str(f.get("name", ""))).suffix.lower().lstrip(".")
    return "image" if ext in _IMAGE else "audio" if ext in _AUDIO else None


def _media_blocks(files: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Image and audio attachments as Strands content blocks."""
    blocks = []
    for f in files:
        ext = Path(str(f["name"])).suffix.lower().lstrip(".")
        data = Path(str(f["path"])).read_bytes()
        if ext in _IMAGE:
            blocks.append({"image": {"format": _IMAGE[ext], "source": {"bytes": data}}})
        elif ext in _AUDIO:
            blocks.append({"audio": {"format": _AUDIO[ext], "source": {"bytes": data}}})
    return blocks


#: `@name`, or `@"a name with spaces"` -- what the composer inserts.
_MENTION = re.compile(r'(?<![\w@])@(?:"([^"\n]+)"|([\w.-]+))')


#: A PDF given as a link in the message: on the web, or a local file:// path.
_PDF_LINK = re.compile(r"""(?:https?|file)://[^\s<>"')\]]+?\.pdf(?:\?[^\s<>"')\]]*)?(?=[\s<>"')\],.;!?]|$)""", re.I)


def _link_guidance(user_text: str, settings: Settings) -> str:
    """The exact command for each PDF link, as attachments get one.

    Without it the model tried the `read` tool on the URL, which cannot fetch,
    and asked the user to upload the file instead.
    """
    lines = []
    for link in dict.fromkeys(_PDF_LINK.findall(user_text)):
        from urllib.parse import urlparse

        if command := reader_command(settings, urlparse(link).path, link):  # the path names the type, not ?query
            lines.append(f"- {link}: read it with: {command}")
    return ("PDF links in this message (the reader downloads them into your workspace):\n"
            + "\n".join(lines)) if lines else ""


def _mention_guidance(user_text: str, settings: Settings, mcp: Any,
                      workspace: Path | None = None) -> str:
    """`@name` in a message, resolved against workspace files, MCP servers,
    tools and skills.

    The composer's @ menu inserts these. A small model needs the instruction
    spelled out: asked to "use @fetch" it would otherwise still pick whichever
    tool looked closest. Resolved at run time, so a retry or an edit of the
    message gets the same instruction. Unknown names are left alone -- an
    email address or a handle is not a mention.
    """
    names = list(dict.fromkeys(q or n for q, n in _MENTION.findall(user_text or "")))
    if not names:
        return ""
    skills = {s["folder"]: s["name"] for s in SkillManager(settings).enabled()}
    skills.update({s: s for s in list(skills.values())})
    servers = {s["name"]: [t["agent_name"] for t in s["tools"] if t["enabled"]]
               for s in mcp.describe() if s["status"] == "connected"}
    tools = set(ToolManager(settings).builtin_tools())
    if settings.tools.documents:
        tools |= {"create_document", "edit_document"}
    for server_tools in servers.values():
        tools |= set(server_tools)

    files = _scan(workspace) if workspace is not None else {}
    lines = []
    for name in names:
        if name in files:
            shown = model_path(str(workspace / name))  # type: ignore[operator]
            command = reader_command(settings, name, name)
            lines.append(f"- `{name}` is the file at {shown}"
                         + (f"; read it with: {command}" if command else "."))
        elif name in servers:
            lines.append(f"- Use the tools of the MCP server `{name}` "
                         f"({', '.join(servers[name]) or 'none enabled'}), not built-in tools that "
                         "do something similar.")
        elif name in tools:
            lines.append(f"- Use the `{name}` tool for this.")
        elif name in skills:
            lines.append(f"- Use the `{skills[name]}` skill: load it with the `skills` tool first.")
    if not lines:
        return ""
    return "The user picked these with @ in their message:\n" + "\n".join(lines)


def _scan(root: Path) -> dict[str, tuple[int, int]]:
    """`relative/path -> (size, mtime_ns)` for every file in a workspace."""
    out: dict[str, tuple[int, int]] = {}
    if not root.is_dir():
        return out
    for p in root.rglob("*"):
        # Dot-files and dot-folders are Atlas's own bookkeeping (temp files,
        # images a reader extracted): not the user's files, not download chips.
        if p.is_file() and not any(part.startswith(".") for part in p.relative_to(root).parts):
            st = p.stat()
            out[p.relative_to(root).as_posix()] = (st.st_size, st.st_mtime_ns)
    return out


def _changed(before: dict[str, tuple[int, int]], after: dict[str, tuple[int, int]]) -> list[dict[str, Any]]:
    """Files a run created or modified, for the download chips under its answer."""
    out = []
    for name, (size, mtime) in sorted(after.items()):
        if name not in before:
            out.append({"name": name, "size": size, "change": "created"})
        elif before[name] != (size, mtime):
            out.append({"name": name, "size": size, "change": "modified"})
    return out


def _sheet_preview(path: Path, max_rows: int = 500, max_cols: int = 50) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            rows, truncated = [], False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    truncated = True
                    break
                rows.append(["" if v is None else str(v) for v in row[:max_cols]])
            sheets.append({"name": ws.title, "rows": rows, "truncated": truncated})
    finally:
        wb.close()
    return {"kind": "sheet", "sheets": sheets}


def _turn_view(turn: dict[str, Any], versions: list[str]) -> dict[str, Any]:
    return {
        "id": turn["id"],
        "parent_id": turn["parent_turn_id"],
        "user_text": turn["user_text"],
        "files": json.loads(turn["files_json"] or "[]"),
        "refs": json.loads(turn.get("refs_json") or "[]"),
        "status": turn["status"],
        "created_at": turn["created_at"],
        "render": json.loads(turn["render_json"]) if turn["render_json"] else [],
        "stats": json.loads(turn["stats_json"]) if turn["stats_json"] else None,
        "error": turn["error"],
        "version": versions.index(turn["id"]) if turn["id"] in versions else 0,
        "versions": versions,
    }


def _run_view(run: Run) -> dict[str, Any]:
    return {"id": run.id, "chat_id": run.chat_id, "turn_id": run.turn_id, "status": run.status,
            "tokens_per_second": run.transcript.stats.get("tokens_per_second")}


def _agent_text(user_text: str, files: list[dict[str, Any]],
                refs: list[dict[str, Any]] | None = None,
                reader: Any = None, workspace: Path | None = None) -> str:
    """What the agent is told, as opposed to what the user typed.

    Nothing is parsed at upload, but a file the model has never heard of is a
    file it will not look for. So attachments are named explicitly, as
    absolute paths. Built here rather than stored, so the stored `user_text`
    stays exactly what was typed -- which is what the transcript shows and what
    a retry replays.
    """
    parts = [user_text]
    quoted = [r for r in refs or [] if isinstance(r, dict) and r.get("quote")]
    if quoted:
        # Selected in the preview panel. Named with the file's path, so a
        # request like "reword this" can go straight to edit_document.
        parts.append("The user selected these passages and is asking about them:\n" + "\n".join(
            f"- from {r.get('file')} at {model_path(r['path']) if r.get('path') else r.get('file')}: "
            f'"{r["quote"]}"' for r in quoted))
    lines = []
    for f in files:
        if not isinstance(f, dict) or not f.get("path"):
            continue
        shown = model_path(f["path"])
        line = f"- {f.get('name')} at {shown} ({f.get('size', 0)} bytes)"
        # The exact command for a pdf/docx/xlsx: the model runs it directly
        # instead of trying `read` first and loading the skill to find out.
        # Commands start in the workspace, so a file there is named bare.
        here = workspace is not None and Path(str(f["path"])).parent == workspace
        if reader is not None and (command := reader(str(f["path"]), f["name"] if here else shown)):
            line += f". Read it with: {command}"
            # Without this the model read the file and rewrote it with
            # create_document even when asked to merge or convert it.
            skill = command.split("$ATLAS_SKILLS/", 1)[-1].split("/", 1)[0]
            line +=f" (to merge, split, combine or convert it instead, load the {skill} skill first)"
        lines.append(line)
    if lines:
        parts.append("Files attached to this message (use the paths exactly as written):\n"
                     + "\n".join(lines))
    return "\n\n".join(p for p in parts if p).strip()


def _strip_hint(messages: list[dict[str, Any]], start: int, hint: str) -> None:
    """Remove a retry hint from the user message before the snapshot is saved.

    The hint is for this one attempt. Left in history it would be re-read on
    every later turn and tell the model to "answer differently" forever.
    """
    if start >= len(messages):
        return
    for block in messages[start].get("content") or []:
        if isinstance(block, dict) and isinstance(block.get("text"), str):
            block["text"] = block["text"].replace(f"\n\n{hint}", "")


__all__ = ["Atlas", "create_app"]
